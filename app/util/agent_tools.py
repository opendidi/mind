# -*- coding: UTF-8 -*-
"""Agent tools for mind — canvas operations, blueprint management, layout, file search, code generation."""

import json
import logging
import os
import re
import tempfile
import uuid
from html.parser import HTMLParser

import requests

from app.util.tool_registry import ToolRegistry
from app.util.vision import VisionHandler
from app.package.module.blueprint_mysql import BlueprintMysqlHandler

# ══════════════════════════════════════════════════════════════════════════════
# Canvas Tools — unified single tool with action parameter
# ══════════════════════════════════════════════════════════════════════════════

_pen_counter = 0


def _next_pen_id():
    global _pen_counter
    _pen_counter += 1
    return f"pen_{_pen_counter}"


@ToolRegistry.register(
    "canvas",
    "画布操作。action: add_pen(创建图形), add_line(连线), add_diagram(批量生成图表), "
    "update_pen(修改), delete_pen(删除), clear(清空), undo(撤销), redo(重做), get_state(查看状态)",
    {"type": "object", "properties": {
        "action": {"type": "string", "enum": [
            "add_pen", "add_line", "add_diagram",
            "update_pen", "delete_pen",
            "clear", "undo", "redo", "get_state",
        ], "description": "操作类型。add_diagram用于批量创建完整图表(流程图/架构图/思维导图)"},
        # add_pen params
        "type": {"type": "string", "description": "[add_pen/add_diagram.nodes]图形类型:rectangle/circle/triangle/diamond/pentagon/star/text/image"},
        "text": {"type": "string", "description": "[add_pen/add_line/update_pen]文字"},
        "x": {"type": "number", "description": "[add_pen]X坐标(画布中心为原点)", "default": 0},
        "y": {"type": "number", "description": "[add_pen]Y坐标(画布中心为原点)", "default": 0},
        "width": {"type": "number", "description": "[add_pen]宽度(像素)", "default": 100},
        "height": {"type": "number", "description": "[add_pen]高度(像素)", "default": 60},
        "background": {"type": "string", "description": "[add_pen]背景颜色(#RRGGBB)"},
        "color": {"type": "string", "description": "[add_pen/add_line]颜色(#RRGGBB)"},
        "fontSize": {"type": "number", "description": "[add_pen]文字大小(px)"},
        "borderWidth": {"type": "number", "description": "[add_pen]边框宽度(px)"},
        "borderColor": {"type": "string", "description": "[add_pen]边框颜色(#RRGGBB)"},
        # add_line params
        "from_pen": {"type": "string", "description": "[add_line/add_diagram.edges]起始节点ID"},
        "to_pen": {"type": "string", "description": "[add_line/add_diagram.edges]目标节点ID"},
        "line_type": {"type": "string", "enum": ["straight", "curve", "polyline", "mind"], "description": "[add_line]连线类型", "default": "straight"},
        "arrow": {"type": "string", "enum": ["start", "end", "both", "none"], "description": "[add_line]箭头方向", "default": "end"},
        "lineWidth": {"type": "number", "description": "[add_line]连线宽度(px)"},
        # add_diagram params
        "diagram": {"type": "object", "description": "[add_diagram]图表定义,含nodes数组和edges数组",
            "properties": {
                "nodes": {"type": "array", "items": {"type": "object"}, "description": "节点列表,每个节点有id/type/text/x/y/width/height"},
                "edges": {"type": "array", "items": {"type": "object"}, "description": "连线列表,每条线有from/to/text/line_type/arrow"},
            },
        },
        # update_pen / delete_pen params
        "pen_id": {"type": "string", "description": "[update_pen/delete_pen]图形ID"},
        "pen_ids": {"type": "array", "items": {"type": "string"}, "description": "[delete_pen]批量删除的图形ID列表"},
        "props": {"type": "object", "description": "[update_pen]要修改的属性键值对,如{\"x\":100,\"text\":\"新文字\"}"},
        # clear param
        "confirm": {"type": "boolean", "description": "[clear]确认清空画布"},
    }, "required": ["action"]}
)
def _tool_canvas(args):
    action = args.get("action", "")
    if action == "add_pen":
        pen_id = _next_pen_id()
        pen_type = args.get("type", "rectangle")
        text = args.get("text", "")
        x, y = args.get("x", 0), args.get("y", 0)
        label = f"「{text}」" if text else ""
        return {
            "success": True, "pen_id": pen_id,
            "data": {**args, "pen_id": pen_id},
            "message": f"已创建{pen_type}{label}，位置({x}, {y})，ID: {pen_id}",
        }
    elif action == "add_line":
        label = f"「{args.get('text')}」" if args.get("text") else ""
        return {
            "success": True, "data": args,
            "message": f"已创建从 {args.get('from_pen')} 到 {args.get('to_pen')} 的{args.get('line_type', 'straight')}连线{label}",
        }
    elif action == "update_pen":
        props_desc = ", ".join(f"{k}={v}" for k, v in args.get("props", {}).items())
        return {"success": True, "data": args, "message": f"已更新图形 {args.get('pen_id')}：{props_desc}"}
    elif action == "delete_pen":
        count = len(args.get("pen_ids", [])) or (1 if args.get("pen_id") else 0)
        return {"success": True, "data": args, "message": f"已删除{count}个图形"}
    elif action == "clear":
        if not args.get("confirm"):
            return {"success": False, "error": "请确认清空画布操作"}
        return {"success": True, "data": {}, "message": "画布已清空"}
    elif action == "undo":
        return {"success": True, "data": {}, "message": "已撤销"}
    elif action == "redo":
        return {"success": True, "data": {}, "message": "已重做"}
    elif action == "add_diagram":
        diagram = args.get("diagram", {})
        nodes = diagram.get("nodes", [])
        edges = diagram.get("edges", [])
        if not nodes:
            return {"success": False, "error": "diagram.nodes 不能为空"}
        # Map LLM-assigned logical IDs → backend pen_ids
        id_map = {}
        for node in nodes:
            logical_id = node.get("id", "")
            pen_id = _next_pen_id()
            node["pen_id"] = pen_id
            if logical_id:
                id_map[logical_id] = pen_id
        # Resolve edge from/to references
        for edge in edges:
            from_ref = edge.get("from", "")
            to_ref = edge.get("to", "")
            edge["_from_id"] = id_map.get(from_ref, from_ref)
            edge["_to_id"] = id_map.get(to_ref, to_ref)
        return {
            "success": True,
            "data": {"diagram": {"nodes": nodes, "edges": edges}},
            "message": f"已生成包含 {len(nodes)} 个节点和 {len(edges)} 条连线的图表",
        }
    elif action == "get_state":
        return {
            "success": True,
            "data": {"note": "画布状态已在系统提示的「当前画布状态」中提供，请基于该上下文和之前的工具调用结果了解当前画布内容。"},
            "message": "画布状态详见系统提示中的「当前画布状态」以及此前的工具调用结果。如两者均为空，则可直接开始创建图形。",
        }
    return {"success": False, "error": f"未知的 action: {action}"}


# ══════════════════════════════════════════════════════════════════════════════
# Blueprint Management Tools
# ══════════════════════════════════════════════════════════════════════════════

@ToolRegistry.register(
    "blueprint_list",
    "列出所有已保存的蓝图，支持分页和筛选。",
    {"type": "object", "properties": {
        "page": {"type": "integer", "description": "页码", "default": 1},
        "page_size": {"type": "integer", "description": "每页数量", "default": 20},
        "keyword": {"type": "string", "description": "搜索关键词"},
    }, "required": []}
)
def _tool_blueprint_list(args):
    user_id = args.get("user_id", "")
    if not user_id:
        return {"success": False, "error": "缺少用户ID，无法查询图纸"}
    data = BlueprintMysqlHandler.query_list({
        'current': args.get('page', 1),
        'page_size': args.get('page_size', 20),
        'keyword': args.get('keyword'),
    }, user_id)
    if data:
        return {"success": True, "data": data, "message": f"共 {data.get('total', 0)} 条图纸"}
    return {"success": True, "data": {"list": [], "total": 0}, "message": "暂无图纸"}


@ToolRegistry.register(
    "blueprint_load",
    "加载指定的蓝图到画布上，替换当前画布内容。",
    {"type": "object", "properties": {
        "blueprint_id": {"type": "string", "description": "要加载的蓝图ID"},
    }, "required": ["blueprint_id"]}
)
def _tool_blueprint_load(args):
    user_id = args.get("user_id", "")
    blueprint_id = args.get("blueprint_id")
    if not blueprint_id:
        return {"success": False, "error": "缺少 blueprint_id"}
    data = BlueprintMysqlHandler.find(blueprint_id, user_id)
    if data:
        return {"success": True, "data": data, "message": f"已加载图纸: {data.get('name', '')}"}
    return {"success": False, "error": f"未找到图纸: {blueprint_id}"}


@ToolRegistry.register(
    "blueprint_save",
    "将当前画布内容保存为蓝图。",
    {"type": "object", "properties": {
        "name": {"type": "string", "description": "蓝图名称"},
        "description": {"type": "string", "description": "蓝图描述"},
        "category": {"type": "string", "description": "分类"},
    }, "required": ["name"]}
)
def _tool_blueprint_save(args):
    user_id = args.get("user_id", "")
    name = args.get("name", "").strip()
    if not name:
        return {"success": False, "error": "图纸名称不能为空"}
    ok, result = BlueprintMysqlHandler.add({
        'name': name,
        'color': args.get('color', ''),
        'penBackground': args.get('penBackground', ''),
        'background': args.get('background', ''),
        'bkImage': args.get('bkImage', ''),
        'grid': args.get('grid', ''),
        'gridColor': args.get('gridColor', ''),
        'gridSize': args.get('gridSize', ''),
        'gridRotate': args.get('gridRotate', ''),
        'rule': args.get('rule', ''),
        'ruleColor': args.get('ruleColor', ''),
        'initJs': args.get('initJs', ''),
        'pens': args.get('pens', ''),
        'https': args.get('https', ''),
        'thumbnail': args.get('thumbnail', ''),
    }, user_id)
    if ok:
        return {"success": True, "data": {"id": result, "name": name}, "message": f"已保存图纸「{name}」"}
    return {"success": False, "error": f"保存失败: {result}"}


@ToolRegistry.register(
    "blueprint_search",
    "按关键词搜索蓝图。",
    {"type": "object", "properties": {
        "keyword": {"type": "string", "description": "搜索关键词"},
        "limit": {"type": "integer", "description": "返回数量上限", "default": 10},
    }, "required": ["keyword"]}
)
def _tool_blueprint_search(args):
    user_id = args.get("user_id", "")
    keyword = args.get("keyword")
    if not keyword:
        return {"success": False, "error": "缺少搜索关键词"}
    data = BlueprintMysqlHandler.query_list({
        'current': 1,
        'page_size': args.get('limit', 10),
        'keyword': keyword,
    }, user_id)
    if data:
        return {"success": True, "data": data, "message": f"搜索「{keyword}」找到 {data.get('total', 0)} 条结果"}
    return {"success": True, "data": {"list": [], "total": 0}, "message": f"搜索「{keyword}」无结果"}


@ToolRegistry.register(
    "blueprint_export",
    "将当前画布导出为指定格式。",
    {"type": "object", "properties": {
        "format": {"type": "string", "enum": ["png", "svg", "json"], "description": "导出格式"},
        "scale": {"type": "number", "description": "缩放比例(PNG导出)", "default": 2},
    }, "required": ["format"]}
)
def _tool_blueprint_export(args):
    return {"success": True, "data": {"format": args.get("format")}, "message": f"已导出为 {args.get('format').upper()} 格式"}


# ══════════════════════════════════════════════════════════════════════════════
# Layout Tools
# ══════════════════════════════════════════════════════════════════════════════

@ToolRegistry.register(
    "layout_auto_arrange",
    "对画布上选中的图形进行自动排版布局。",
    {"type": "object", "properties": {
        "direction": {"type": "string", "enum": ["horizontal", "vertical", "grid"], "description": "排列方向", "default": "vertical"},
        "spacing": {"type": "number", "description": "间距(px)", "default": 40},
        "columns": {"type": "integer", "description": "网格列数(grid模式)", "default": 3},
    }, "required": []}
)
def _tool_layout_auto_arrange(args):
    return {"success": True, "data": args, "message": f"已按{args.get('direction', 'vertical')}方向自动排列，间距{args.get('spacing', 40)}px"}


@ToolRegistry.register(
    "layout_align",
    "将选中的多个图形对齐。",
    {"type": "object", "properties": {
        "align": {"type": "string", "enum": ["left", "center", "right", "top", "middle", "bottom"], "description": "对齐方式"},
    }, "required": ["align"]}
)
def _tool_layout_align(args):
    return {"success": True, "data": args, "message": f"已按{args.get('align')}对齐"}


# ══════════════════════════════════════════════════════════════════════════════
# Auxiliary Tools
# ══════════════════════════════════════════════════════════════════════════════

@ToolRegistry.register(
    "file_search",
    "搜索文件管理器中的素材/文件。",
    {"type": "object", "properties": {
        "keyword": {"type": "string", "description": "文件名关键词"},
        "type": {"type": "string", "description": "文件类型筛选: image/svg/document"},
        "limit": {"type": "integer", "description": "返回数量", "default": 20},
    }, "required": []}
)
def _tool_file_search(args):
    return {"success": True, "data": {"items": []}, "message": "文件搜索结果（需连接存储）"}


@ToolRegistry.register(
    "code_generate",
    "根据描述生成JavaScript或JSON代码片段，可用于Monaco编辑器中。",
    {"type": "object", "properties": {
        "language": {"type": "string", "enum": ["javascript", "json"], "description": "代码语言"},
        "description": {"type": "string", "description": "代码需求描述"},
    }, "required": ["language", "description"]}
)
def _tool_code_generate(args):
    return {"success": True, "data": {"language": args.get("language"), "code": "// Generated code"}, "message": "代码已生成"}


# ══════════════════════════════════════════════════════════════════════════════
# Search & Web Tools
# ══════════════════════════════════════════════════════════════════════════════

def _require(args: dict, *keys: str) -> str | None:
    for k in keys:
        val = args.get(k)
        if val is None or (isinstance(val, str) and not val.strip()):
            return f"缺少必填参数: {k}"
    return None


class _TextExtractor(HTMLParser):
    """Extract plain text from HTML, stripping all tags, skipping scripts/styles."""

    def __init__(self):
        super().__init__()
        self.text = []
        self._skip = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript", "iframe", "head"):
            self._skip = True

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript", "iframe", "head"):
            self._skip = False
        if tag in ("p", "br", "div", "li", "h1", "h2", "h3", "h4", "h5", "h6", "tr"):
            self.text.append("\n")

    def handle_data(self, data):
        if not self._skip:
            t = data.strip()
            if t:
                self.text.append(t)

    def get_text(self):
        return "\n".join(self.text)


# web_search tool is now in app.util.search.engine_chain (auto-registered on import)
# web_fetch tool kept below

@ToolRegistry.register(
    "web_fetch",
    "抓取指定URL的网页内容并提取正文。适用于用户提供链接要求分析、总结或提取页面信息时使用。",
    {"type": "object", "properties": {
        "url": {"type": "string", "description": "要抓取的网页链接（完整URL，如 https://example.com/article）"},
    }, "required": ["url"]},
)
def _tool_web_fetch(args):
    url = args.get("url", "").strip()
    if not url:
        return {"success": False, "error": "请提供要抓取的链接"}
    if not url.startswith(("http://", "https://")):
        url = "https://" + url

    try:
        resp = requests.get(
            url,
            timeout=15,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/120.0.0.0 Safari/537.36"
                ),
            },
        )
        resp.raise_for_status()

        ct = resp.headers.get("Content-Type", "")
        if "text/html" not in ct.lower() and "text/plain" not in ct.lower():
            return {
                "success": False,
                "error": f"该链接不是HTML页面（Content-Type: {ct}），无法提取文本",
            }

        resp.encoding = resp.apparent_encoding or "utf-8"
        html_text = resp.text

        title_match = re.search(
            r"<title[^>]*>(.*?)</title>", html_text, re.IGNORECASE | re.DOTALL
        )
        title = title_match.group(1).strip() if title_match else ""

        extractor = _TextExtractor()
        extractor.feed(html_text)
        body = extractor.get_text()

        if len(body) > 8000:
            body = body[:8000] + "\n\n... [内容已截断，原文过长]"

        return {
            "success": True,
            "data": {
                "url": url,
                "title": title[:200],
                "content": body,
                "length": len(body),
            },
        }
    except requests.exceptions.Timeout:
        return {"success": False, "error": f"请求超时 (15s): {url}"}
    except requests.exceptions.HTTPError as e:
        return {"success": False, "error": f"请求失败 (HTTP {e.response.status_code}): {url}"}
    except requests.exceptions.ConnectionError:
        return {"success": False, "error": f"无法连接到该网站: {url}"}
    except Exception as e:
        logging.warning("web_fetch 失败: %s", e)
        return {"success": False, "error": f"抓取失败: {str(e)[:150]}"}


# ══════════════════════════════════════════════════════════════════════════════
# File Analysis Tools
# ══════════════════════════════════════════════════════════════════════════════

def _download_from_minio(object_name: str) -> str | None:
    """Download a file from MinIO to a temp path. Returns the local path or None."""
    from app.plugin.minio.app.controller import minio_client, bucket_name

    _, ext = os.path.splitext(object_name)
    fd, tmp_path = tempfile.mkstemp(suffix=ext)
    os.close(fd)
    try:
        minio_client.fget_object(bucket_name, object_name, tmp_path)
        return tmp_path
    except Exception:
        try:
            os.remove(tmp_path)
        except OSError:
            pass
        return None


def _validate_analyze_image(args):
    url = args.get("image_url", "")
    if not url or not isinstance(url, str) or not url.strip():
        return "缺少必填参数: image_url"
    task = args.get("task_type", "")
    if task not in ("describe", "detect", "classify"):
        return f"task_type 无效: {task}，支持 describe/detect/classify"
    return None


@ToolRegistry.register(
    "analyze_image",
    "分析一张图片的内容。可进行场景描述(describe)、物体检测(detect)或智能分类(classify)。",
    {"type": "object", "properties": {
        "image_url": {"type": "string", "description": "图片URL地址"},
        "task_type": {
            "type": "string",
            "enum": ["describe", "detect", "classify"],
            "description": "describe=场景描述(含关键词), detect=物体检测, classify=智能分类",
        },
    }, "required": ["image_url", "task_type"]},
    validator=_validate_analyze_image,
)
def _tool_analyze_image(args):
    image_url = args["image_url"]
    task_type = args["task_type"]

    if task_type == "classify":
        prompt = VisionHandler.build_classify_prompt([])
    elif task_type == "detect":
        prompt = VisionHandler.build_detect_prompt()
    else:
        prompt = VisionHandler.build_describe_prompt()

    ok, result = VisionHandler.analyze_image(image_url, prompt, task_type)
    return {"success": ok, "data": result, "task_type": task_type}


def _validate_analyze_doc(args):
    obj = args.get("object_name", "")
    if not obj or not isinstance(obj, str) or not obj.strip():
        return "缺少必填参数: object_name"
    if not obj.lower().endswith(".docx"):
        return f"不支持的文件格式，仅支持 .docx: {obj}"
    return None


@ToolRegistry.register(
    "analyze_doc",
    "解析 Word 文档 (.docx)，提取段落文本、标题和表格内容。支持分页读取长文档。",
    {"type": "object", "properties": {
        "object_name": {"type": "string", "description": "MinIO 中的文档对象路径 (object_name)，由文件上传接口返回"},
        "paragraph_offset": {"type": "integer", "description": "段落起始偏移量（默认 0，用于分页读取长文档）"},
        "paragraph_limit": {"type": "integer", "description": "最多返回的段落数（默认 200，最大 500）"},
    }, "required": ["object_name"]},
    validator=_validate_analyze_doc,
)
def _tool_analyze_doc(args):
    object_name = args["object_name"].strip()
    para_offset = max(0, args.get("paragraph_offset", 0) or 0)
    para_limit = min(args.get("paragraph_limit", 200) or 200, 500)

    tmp_path = _download_from_minio(object_name)
    if not tmp_path:
        return {"success": False, "error": f"从 MinIO 下载文件失败: {object_name}"}

    try:
        from docx import Document

        doc = Document(tmp_path)
        paragraphs = []
        tables_data = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                style = p.style.name if p.style else ""
                paragraphs.append({"text": text, "style": style})

        for ti, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            tables_data.append({"table_index": ti, "rows": rows})

        total_paras = len(paragraphs)
        paged = paragraphs[para_offset : para_offset + para_limit]
        truncated = (para_offset + para_limit) < total_paras
        raw_text = "\n\n".join(p["text"] for p in paged)

        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "paragraphs_count": total_paras,
                "tables_count": len(tables_data),
                "paragraph_offset": para_offset,
                "paragraph_limit": para_limit,
                "paragraphs": paged,
                "tables": tables_data[:30],
                "raw_text": raw_text[:40000],
                "truncated": truncated,
                "hint": f"共 {total_paras} 段落，当前返回第 {para_offset}–{para_offset + len(paged)} 段。"
                + ("内容未读完，用 paragraph_offset 参数继续。" if truncated else "已读完。"),
            },
        }
    except Exception as ex:
        logging.warning("analyze_doc 解析失败: %s", ex)
        return {"success": False, "error": f"文档解析失败: {str(ex)}"}
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass


def _validate_extract_excel(args):
    obj = args.get("object_name", "")
    if not obj or not isinstance(obj, str) or not obj.strip():
        return "缺少必填参数: object_name"
    if not (obj.lower().endswith(".xlsx") or obj.lower().endswith(".xls")):
        return f"不支持的文件格式，仅支持 .xlsx/.xls: {obj}"
    return None


@ToolRegistry.register(
    "extract_excel",
    "解析 Excel 文档 (.xlsx/.xls)，读取工作表数据。支持分页读取大表格。",
    {"type": "object", "properties": {
        "object_name": {"type": "string", "description": "MinIO 中的 Excel 对象路径 (object_name)，由文件上传接口返回"},
        "sheet_name": {"type": "string", "description": "要读取的工作表名称（可选，不传则读取第一个工作表）"},
        "row_offset": {"type": "integer", "description": "行起始偏移量（默认 0，用于分页读取大表格）"},
        "max_rows": {"type": "integer", "description": "最多读取的行数（默认 500，最大 2000）"},
    }, "required": ["object_name"]},
    validator=_validate_extract_excel,
)
def _tool_extract_excel(args):
    object_name = args["object_name"].strip()
    sheet_name = (args.get("sheet_name") or "").strip() or None
    row_offset = max(0, args.get("row_offset", 0) or 0)
    max_rows = min(args.get("max_rows", 500) or 500, 2000)

    tmp_path = _download_from_minio(object_name)
    if not tmp_path:
        return {"success": False, "error": f"从 MinIO 下载文件失败: {object_name}"}

    try:
        is_legacy = object_name.lower().endswith(".xls") and not object_name.lower().endswith(".xlsx")

        if is_legacy:
            import xlrd

            wb = xlrd.open_workbook(tmp_path)
            sheet_names = wb.sheet_names()

            target = sheet_name or sheet_names[0]
            try:
                ws = wb.sheet_by_name(target)
            except xlrd.XLRDError:
                return {
                    "success": False,
                    "error": f"工作表 '{target}' 不存在，可用: {', '.join(sheet_names)}",
                }

            total_rows = ws.nrows
            data = []
            for ri in range(min(total_rows, row_offset + max_rows)):
                if ri < row_offset:
                    continue
                data.append([str(ws.cell_value(ri, ci)) if ws.cell_value(ri, ci) != "" else "" for ci in range(ws.ncols)])
        else:
            from openpyxl import load_workbook

            wb = load_workbook(tmp_path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames

            target = sheet_name or sheet_names[0]
            if target not in sheet_names:
                return {
                    "success": False,
                    "error": f"工作表 '{target}' 不存在，可用: {', '.join(sheet_names)}",
                }

            ws = wb[target]
            data = []
            total_rows = 0
            for ri, row in enumerate(ws.iter_rows(values_only=True)):
                total_rows += 1
                if ri < row_offset:
                    continue
                if len(data) >= max_rows:
                    continue
                data.append([str(c) if c is not None else "" for c in row])

            wb.close()

        truncated = (row_offset + len(data)) < total_rows
        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "sheet_name": target,
                "sheet_names": sheet_names,
                "total_rows": total_rows,
                "row_offset": row_offset,
                "max_rows": max_rows,
                "rows_returned": len(data),
                "rows": data,
                "truncated": truncated,
                "hint": f"共 {total_rows} 行，当前返回第 {row_offset}–{row_offset + len(data)} 行。"
                + ("数据未读完，用 row_offset 参数继续。" if truncated else "已读完。"),
            },
        }
    except Exception as ex:
        logging.warning("extract_excel 解析失败: %s", ex)
        return {"success": False, "error": f"Excel 解析失败: {str(ex)}"}
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass


def _validate_read_text(args):
    obj = args.get("object_name", "")
    if not obj or not isinstance(obj, str) or not obj.strip():
        return "缺少必填参数: object_name"
    lower = obj.lower()
    if not (lower.endswith(".txt") or lower.endswith(".md")):
        return f"不支持的文件格式，仅支持 .txt/.md: {obj}"
    return None


@ToolRegistry.register(
    "read_text",
    "读取纯文本/Markdown 文件 (.txt/.md)，支持分页读取长文本。",
    {"type": "object", "properties": {
        "object_name": {"type": "string", "description": "MinIO 中的文本文件 object_name，由文件上传接口返回"},
        "char_offset": {"type": "integer", "description": "字符起始偏移量（默认 0，用于分页读取长文本）"},
        "char_limit": {"type": "integer", "description": "最多返回的字符数（默认 30000，最大 80000）"},
    }, "required": ["object_name"]},
    validator=_validate_read_text,
)
def _tool_read_text(args):
    object_name = args["object_name"].strip()
    char_offset = max(0, args.get("char_offset", 0) or 0)
    char_limit = min(args.get("char_limit", 30000) or 30000, 80000)

    tmp_path = _download_from_minio(object_name)
    if not tmp_path:
        return {"success": False, "error": f"从 MinIO 下载文件失败: {object_name}"}

    def _read(encoding):
        with open(tmp_path, "r", encoding=encoding) as f:
            return f.read()

    try:
        content = _read("utf-8")
    except UnicodeDecodeError:
        try:
            content = _read("gbk")
        except Exception as ex:
            return {"success": False, "error": f"文件编码不支持 (UTF-8/GBK): {str(ex)}"}
    except Exception as ex:
        logging.warning("read_text 读取失败: %s", ex)
        return {"success": False, "error": f"文件读取失败: {str(ex)}"}
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass

    total = len(content)
    chunk = content[char_offset : char_offset + char_limit]
    truncated = (char_offset + char_limit) < total
    return {
        "success": True,
        "data": {
            "filename": os.path.basename(object_name),
            "content": chunk,
            "total_chars": total,
            "char_offset": char_offset,
            "char_limit": char_limit,
            "truncated": truncated,
            "hint": f"共 {total} 字符，当前返回第 {char_offset}–{char_offset + len(chunk)} 字符。"
            + ("内容未读完，用 char_offset 参数继续。" if truncated else "已读完。"),
        },
    }


def _validate_parse_json(args):
    obj = args.get("object_name", "")
    if not obj or not isinstance(obj, str) or not obj.strip():
        return "缺少必填参数: object_name"
    if not obj.lower().endswith(".json"):
        return f"不支持的文件格式，仅支持 .json: {obj}"
    return None


@ToolRegistry.register(
    "parse_json",
    "解析 JSON 文件 (.json)，返回结构化数据摘要。支持数组/对象分页、路径导航、大值截断。",
    {"type": "object", "properties": {
        "object_name": {"type": "string", "description": "MinIO 中的 JSON 文件 object_name，由文件上传接口返回"},
        "query_path": {"type": "string", "description": "可选的路径（如 'data.users'、'items[0].name'），用于提取深层字段"},
        "array_offset": {"type": "integer", "description": "数组元素起始索引（默认 0）"},
        "array_limit": {"type": "integer", "description": "数组最多返回元素数（默认 200，最大 2000）"},
        "key_offset": {"type": "integer", "description": "对象键起始偏移量（默认 0）"},
        "key_limit": {"type": "integer", "description": "对象最多返回键数（默认 50，最大 200）"},
        "max_list_items": {"type": "integer", "description": "嵌套列表最多返回项数（默认 100，最大 2000）"},
    }, "required": ["object_name"]},
    validator=_validate_parse_json,
)
def _tool_parse_json(args):
    object_name = args["object_name"].strip()
    query_path = (args.get("query_path") or "").strip() or None
    array_offset = max(0, args.get("array_offset", 0) or 0)
    array_limit = min(args.get("array_limit", 200) or 200, 2000)
    key_offset = max(0, args.get("key_offset", 0) or 0)
    key_limit = min(args.get("key_limit", 50) or 50, 200)
    max_list = min(args.get("max_list_items", 100) or 100, 2000)

    MAX_VAL_LEN = 8000
    MAX_DICT_SAMPLE = 20
    MAX_DEPTH = 4

    def _summarize(val, depth=0, field_path=""):
        if depth > MAX_DEPTH:
            return f"<嵌套{type(val).__name__}>"
        if isinstance(val, str):
            if len(val) > MAX_VAL_LEN:
                return val[:MAX_VAL_LEN] + f"…(截断，共{len(val)}字符)"
            return val
        if isinstance(val, (int, float, bool)) or val is None:
            return val
        if isinstance(val, list):
            total = len(val)
            if total == 0:
                return []
            if total <= max_list:
                return [_summarize(v, depth + 1, f"{field_path}[{i}]") for i, v in enumerate(val)]
            sample = [_summarize(v, depth + 1, f"{field_path}[{i}]") for i, v in enumerate(val[:max_list])]
            return {
                "__head": sample,
                "__total": total,
                "__hint": f"数组共 {total} 项，仅返回前 {max_list} 项。用 query_path 定位到此字段后可用 array_offset/array_limit 分页读完。",
            }
        if isinstance(val, dict):
            ks = list(val.keys())
            total = len(ks)
            if total == 0:
                return {}
            if total <= MAX_DICT_SAMPLE:
                return {k: _summarize(v, depth + 1, f"{field_path}.{k}" if field_path else k)
                        for k, v in val.items()}
            head = {k: _summarize(v, depth + 1, f"{field_path}.{k}" if field_path else k)
                    for k, v in list(val.items())[:MAX_DICT_SAMPLE]}
            return {"__dict_head": head, "__total_keys": total}
        return str(val)[:500]

    tmp_path = _download_from_minio(object_name)
    if not tmp_path:
        return {"success": False, "error": f"从 MinIO 下载文件失败: {object_name}"}

    try:
        with open(tmp_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except UnicodeDecodeError:
        try:
            with open(tmp_path, "r", encoding="gbk") as f:
                data = json.load(f)
        except Exception as ex:
            return {"success": False, "error": f"JSON 解析失败 (编码): {str(ex)}"}
    except json.JSONDecodeError as ex:
        return {"success": False, "error": f"JSON 格式无效: {str(ex)}"}
    except Exception as ex:
        logging.warning("parse_json 读取失败: %s", ex)
        return {"success": False, "error": f"文件读取失败: {str(ex)}"}
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass

    # Path navigation
    target = data
    if query_path:
        try:
            for key in query_path.split("."):
                if "[" in key and key.endswith("]"):
                    name, idx = key[:-1].split("[", 1)
                    target = target[name][int(idx)]
                else:
                    target = target[key]
        except (KeyError, IndexError, TypeError, ValueError) as ex:
            return {"success": False, "error": f"路径 '{query_path}' 不可达: {str(ex)}"}

    if isinstance(target, list):
        total = len(target)
        page = target[array_offset : array_offset + array_limit]
        truncated = (array_offset + array_limit) < total
        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "type": "array",
                "total_items": total,
                "array_offset": array_offset,
                "array_limit": array_limit,
                "items_returned": len(page),
                "items": [_summarize(v) for v in page],
                "truncated": truncated,
                "hint": f"JSON 数组共 {total} 项，当前返回第 {array_offset}–{array_offset + len(page)} 项。"
                + ("未读完，用 array_offset 参数继续。" if truncated else "已读完。"),
            },
        }
    elif isinstance(target, dict):
        all_keys = list(target.keys())
        total_keys = len(all_keys)
        page_keys = all_keys[key_offset : key_offset + key_limit]
        page_items = {k: _summarize(target[k]) for k in page_keys}
        truncated = (key_offset + key_limit) < total_keys
        raw_size = len(json.dumps(target, ensure_ascii=False))
        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "type": "object",
                "total_keys": total_keys,
                "key_offset": key_offset,
                "key_limit": key_limit,
                "keys_returned": len(page_keys),
                "items": page_items,
                "raw_size": raw_size,
                "truncated": truncated,
                "hint": f"JSON 对象共 {total_keys} 个键（{raw_size} 字符），当前返回第 {key_offset}–{key_offset + len(page_keys)} 个键。"
                + ("未读完，用 key_offset 参数继续。" if truncated else "已读完。"),
            },
        }
    else:
        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "type": type(target).__name__,
                "value": _summarize(target),
            },
        }


# ══════════════════════════════════════════════════════════════════════════════
# Map / Geo Tools — geocode, regeocode
# ══════════════════════════════════════════════════════════════════════════════

from app.config import AMAP_KEY


def _validate_geocode(args):
    addr = args.get("address")
    if not addr or not isinstance(addr, str) or not addr.strip():
        return "address 必须是非空字符串"
    return None


@ToolRegistry.register(
    "geocode",
    "将地名/地址转换为经纬度坐标。用于路线规划前获取准确的起点/终点坐标。"
    "例如用户说'从北京到上海'，先调用此工具查询'北京'和'上海'获取坐标。"
    "返回：匹配地点列表，每个含坐标(lng,lat)、地址名称、行政区划。",
    {
        "type": "object",
        "properties": {
            "address": {"type": "string", "description": "要查询的地点名称或地址，如'广州塔'、'北京市朝阳区'"},
            "city": {"type": "string", "description": "限定城市范围（可选），如'北京'、'广州'"},
        },
        "required": ["address"],
    },
    validator=_validate_geocode,
)
def _tool_geocode(args):
    address = args["address"].strip()
    city = args.get("city", "").strip() or None

    if not AMAP_KEY:
        return {"success": False, "error": "AMAP_KEY 未配置，请联系管理员在 .env 中设置"}

    params = {"key": AMAP_KEY, "address": address, "output": "JSON"}
    if city:
        params["city"] = city

    try:
        resp = requests.get("https://restapi.amap.com/v3/geocode/geo", params=params, timeout=10)
        data = resp.json()
        if data.get("status") != "1":
            return {"success": False, "error": f"高德地理编码失败: {data.get('info', '未知错误')}"}

        geocodes = data.get("geocodes") or []
        if not geocodes:
            return {"success": False, "error": f"未找到'{address}'的地理位置，请检查地名是否正确"}

        results = []
        for g in geocodes[:5]:
            loc = g.get("location", "")
            lng, lat = (loc.split(",") + ["0", "0"])[:2]
            results.append({
                "lng": float(lng), "lat": float(lat),
                "address": g.get("formatted_address", address),
                "city": g.get("city", ""), "district": g.get("district", ""),
                "level": g.get("level", ""),
            })
        return {"success": True, "data": {"total": len(results), "locations": results}}

    except requests.RequestException as e:
        return {"success": False, "error": f"地理编码请求失败: {str(e)}"}


def _validate_regeocode(args):
    lng = args.get("lng")
    lat = args.get("lat")
    if lng is None or lat is None:
        return "lng 和 lat 是必填项"
    try:
        float(lng)
        float(lat)
    except (TypeError, ValueError):
        return "lng 和 lat 必须是有效数字"
    return None


@ToolRegistry.register(
    "regeocode",
    "将经纬度坐标逆地理编码为地址信息。用于：用户提供坐标想知道是哪里。"
    "返回：格式化地址、省份、城市、区县、附近地标。",
    {
        "type": "object",
        "properties": {
            "lng": {"type": "number", "description": "经度（十进制），如 113.175"},
            "lat": {"type": "number", "description": "纬度（十进制），如 22.793"},
        },
        "required": ["lng", "lat"],
    },
    validator=_validate_regeocode,
)
def _tool_regeocode(args):
    lng = float(args["lng"])
    lat = float(args["lat"])

    if not AMAP_KEY:
        return {"success": False, "error": "AMAP_KEY 未配置，请联系管理员在 .env 中设置"}

    params = {"key": AMAP_KEY, "location": f"{lng},{lat}", "output": "JSON", "extensions": "base"}

    try:
        resp = requests.get("https://restapi.amap.com/v3/geocode/regeo", params=params, timeout=10)
        data = resp.json()
        if data.get("status") != "1":
            return {"success": False, "error": f"高德逆地理编码失败: {data.get('info', '未知错误')}"}

        regeocode = data.get("regeocode") or {}
        addr = regeocode.get("addressComponent") or {}
        return {
            "success": True,
            "data": {
                "address": regeocode.get("formatted_address", ""),
                "country": addr.get("country", ""),
                "province": addr.get("province", ""),
                "city": addr.get("city", ""),
                "district": addr.get("district", ""),
                "township": addr.get("township", ""),
                "road": addr.get("streetNumber", {}).get("street", ""),
                "pois": [{"name": p.get("name"), "type": p.get("type")}
                         for p in (regeocode.get("pois") or [])[:5]],
            },
        }
    except requests.RequestException as e:
        return {"success": False, "error": f"逆地理编码请求失败: {str(e)}"}


# ══════════════════════════════════════════════════════════════════════════════
# Populate TOOL_SCHEMAS and run_tool_call
# ══════════════════════════════════════════════════════════════════════════════

TOOL_SCHEMAS = []  # populated below after all registrations


def _rebuild_schemas():
    """Rebuild TOOL_SCHEMAS from ToolRegistry."""
    global TOOL_SCHEMAS
    TOOL_SCHEMAS = ToolRegistry.get_schemas()


_rebuild_schemas()


def run_tool_call(tool_name, tool_args, tool_context, dispatcher=None,
                  llm_client=None, model=None, tracer=None, pheromone_sniff="",
                  event_queue=None):
    """Execute a tool by name, dispatching sub-agents when needed.

    Returns (result_dict, cached_bool).
    """
    # Handle dispatch_agent — route to sub-agent system
    if tool_name == "dispatch_agent" and dispatcher:
        agent_name = tool_args.get("agent_name", "")
        task = tool_args.get("task", "")
        if agent_name and task:
            ctx = {**tool_context, "_pheromone": pheromone_sniff}
            result = dispatcher.dispatch(llm_client, agent_name, task, ctx, model=model, tracer=tracer,
                                        event_queue=event_queue, stream=(event_queue is not None))
            return result, False

    # Handle ask_peer — lightweight sub-agent query
    if tool_name == "ask_peer" and dispatcher:
        agent_name = tool_args.get("agent_name", "")
        question = tool_args.get("question", "")
        if agent_name and question:
            ctx = {**tool_context, "_pheromone": pheromone_sniff}
            result = dispatcher.handle_peer_query(llm_client, agent_name, question, ctx, model=model, tracer=tracer)
            return result, False

    # Normal tool execution via ToolRegistry
    result = ToolRegistry.execute(tool_name, tool_args, tool_context)
    return result, False

