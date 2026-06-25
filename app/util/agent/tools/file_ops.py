# -*- coding: UTF-8 -*-
"""File operations tool — search, download, and analyze files."""

import json
import logging
import os
import re
import tempfile
import uuid
from html.parser import HTMLParser

import requests

from app.package.module.blueprint_mysql import BlueprintMysqlHandler
from app.util.tool_registry import ToolRegistry
from app.util.vision import VisionHandler

# ══════════════════════════════════════════════════════════════════════════════
# Auxiliary Tools
# ══════════════════════════════════════════════════════════════════════════════


@ToolRegistry.register(
    "file_search",
    "搜索文件管理器中的素材/文件。",
    {
        "type": "object",
        "properties": {
            "keyword": {"type": "string", "description": "文件名关键词"},
            "type": {"type": "string", "description": "文件类型筛选: image/svg/document/text"},
            "limit": {"type": "integer", "description": "返回数量", "default": 20},
        },
        "required": [],
    },
)
def _tool_file_search(args):
    """Search the file manager material table for uploaded files."""
    user_id = args.get("user_id", "")
    keyword = args.get("keyword", "").strip()
    file_type = args.get("type") or None
    limit = max(1, min(int(args.get("limit", 20)), 50))

    try:
        from app.package.module.material_mysql import MaterialMysqlHandler

        result = MaterialMysqlHandler.query_list(
            {
                "current": 1,
                "page_size": limit,
                "keyword": keyword or None,
                "type": file_type,
                "folder": None,
                "parent_id": None,
                "sort_order": "created_at DESC",
            },
            user_id,
        )
        # query_list returns {"list": [...], "total": N} — not "data"
        items = result.get("list", []) if isinstance(result, dict) else []
        if not items:
            return {"success": True, "data": {"items": [], "total": 0}, "message": "没有找到匹配的文件"}
        return {"success": True, "data": {"items": items, "total": result.get("total", len(items))}}
    except Exception as e:
        logging.warning(f"file_search 查询失败：{e}")
        return {"success": False, "message": f"搜索失败：{e}"}


# ══════════════════════════════════════════════════════════════════════════════
# File Analysis Tools
# ══════════════════════════════════════════════════════════════════════════════


def _download_from_minio(object_name: str) -> str | None:
    """Download a file from MinIO to a temp path. Returns the local path or None."""
    from app.plugin.minio.app.controller import bucket_name, minio_client

    _, ext = os.path.splitext(object_name)
    fd, tmp_path = tempfile.mkstemp(suffix=ext)
    os.close(fd)
    try:
        minio_client.fget_object(bucket_name, object_name, tmp_path)
        return tmp_path
    except Exception:
        try:
            os.remove(tmp_path)
        except OSError:
            pass
        return None


def _validate_analyze_image(args):
    url = args.get("image_url", "")
    b64 = args.get("image_base64", "")
    obj = args.get("object_name", "")
    if not (url or b64 or obj):
        return "缺少参数: 需要 image_url / image_base64 / object_name 之一"
    if url and not isinstance(url, str):
        return "image_url 必须是字符串"
    if b64 and not (isinstance(b64, str) and b64.startswith("data:")):
        return "image_base64 必须是 data:image/...;base64,... 格式"
    task = args.get("task_type", "")
    if task not in ("describe", "detect", "classify", "ocr"):
        return f"task_type 无效: {task}，支持 describe/detect/classify/ocr"
    return None


@ToolRegistry.register(
    "analyze_image",
    "分析一张图片的内容。可进行场景描述(describe)、物体检测(detect)、智能分类(classify)或文字提取(ocr)。"
    "三种输入方式任选其一：1) image_url — 图片URL地址（可使用 file_search 返回结果中的 url 字段）；"
    "2) image_base64 — base64格式的图片数据；3) object_name — MinIO 中的图片对象路径（由文件上传接口返回）。",
    {
        "type": "object",
        "properties": {
            "image_url": {
                "type": "string",
                "description": "图片URL地址，可直接使用 file_search 返回结果中的 url 字段（与 image_base64 / object_name 三选一）",
            },
            "image_base64": {
                "type": "string",
                "description": "base64格式的图片数据，data:image/...;base64,... 格式（与 image_url / object_name 三选一）",
            },
            "object_name": {
                "type": "string",
                "description": "MinIO 中的图片对象路径，由文件上传接口返回，或从 file_search 结果的 url 中提取路径（与 image_url / image_base64 三选一）",
            },
            "task_type": {
                "type": "string",
                "enum": ["describe", "detect", "classify", "ocr"],
                "description": "describe=场景描述(含关键词), detect=物体检测, classify=智能分类, ocr=文字提取",
            },
        },
        "required": ["task_type"],
    },
    validator=_validate_analyze_image,
)
def _tool_analyze_image(args):
    import base64 as _b64

    task_type = args["task_type"]
    image_base64 = args.get("image_base64", "")
    image_url = args.get("image_url", "")
    object_name = args.get("object_name", "")

    if task_type == "classify":
        prompt = VisionHandler.build_classify_prompt([])
    elif task_type == "detect":
        prompt = VisionHandler.build_detect_prompt()
    elif task_type == "ocr":
        prompt = VisionHandler.build_ocr_prompt()
    else:
        prompt = VisionHandler.build_describe_prompt()

    # Resolve object_name → base64 via MinIO download
    if object_name and not image_base64 and not image_url:
        tmp_path = _download_from_minio(object_name)
        if not tmp_path:
            return {"success": False, "error": f"从 MinIO 下载图片失败: {object_name}"}
        try:
            with open(tmp_path, "rb") as f:
                raw = f.read()
            ext = os.path.splitext(object_name)[1].lstrip(".").lower()
            mime = {
                "png": "image/png",
                "jpg": "image/jpeg",
                "jpeg": "image/jpeg",
                "gif": "image/gif",
                "webp": "image/webp",
                "svg": "image/svg+xml",
            }.get(ext, "image/png")
            image_base64 = f"data:{mime};base64,{_b64.b64encode(raw).decode()}"
        except Exception as e:
            return {"success": False, "error": f"读取图片文件失败: {e}"}
        finally:
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    if image_base64:
        ok, result = VisionHandler.analyze_base64(image_base64, prompt, task_type)
    else:
        ok, result = VisionHandler.analyze_image(image_url, prompt, task_type)
    return {"success": ok, "data": result, "task_type": task_type}


def _validate_analyze_doc(args):
    obj = args.get("object_name", "")
    if not obj or not isinstance(obj, str) or not obj.strip():
        return "缺少必填参数: object_name"
    if not obj.lower().endswith(".docx"):
        return f"不支持的文件格式，仅支持 .docx: {obj}"
    return None


@ToolRegistry.register(
    "analyze_doc",
    "解析 Word 文档 (.docx)，提取段落文本、标题和表格内容。支持分页读取长文档。",
    {
        "type": "object",
        "properties": {
            "object_name": {
                "type": "string",
                "description": "MinIO 中的文档对象路径 (object_name)，由文件上传接口返回",
            },
            "paragraph_offset": {"type": "integer", "description": "段落起始偏移量（默认 0，用于分页读取长文档）"},
            "paragraph_limit": {"type": "integer", "description": "最多返回的段落数（默认 200，最大 500）"},
        },
        "required": ["object_name"],
    },
    validator=_validate_analyze_doc,
)
def _tool_analyze_doc(args):
    object_name = args["object_name"].strip()
    para_offset = max(0, args.get("paragraph_offset", 0) or 0)
    para_limit = min(args.get("paragraph_limit", 200) or 200, 500)

    tmp_path = _download_from_minio(object_name)
    if not tmp_path:
        return {"success": False, "error": f"从 MinIO 下载文件失败: {object_name}"}

    try:
        from docx import Document

        doc = Document(tmp_path)
        paragraphs = []
        tables_data = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                style = p.style.name if p.style else ""
                paragraphs.append({"text": text, "style": style})

        for ti, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            tables_data.append({"table_index": ti, "rows": rows})

        total_paras = len(paragraphs)
        paged = paragraphs[para_offset : para_offset + para_limit]
        truncated = (para_offset + para_limit) < total_paras
        raw_text = "\n\n".join(p["text"] for p in paged)

        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "paragraphs_count": total_paras,
                "tables_count": len(tables_data),
                "paragraph_offset": para_offset,
                "paragraph_limit": para_limit,
                "paragraphs": paged,
                "tables": tables_data[:30],
                "raw_text": raw_text[:40000],
                "truncated": truncated,
                "hint": f"共 {total_paras} 段落，当前返回第 {para_offset}–{para_offset + len(paged)} 段。"
                + ("内容未读完，用 paragraph_offset 参数继续。" if truncated else "已读完。"),
            },
        }
    except Exception as ex:
        logging.warning("analyze_doc 解析失败: %s", ex)
        return {"success": False, "error": f"文档解析失败: {str(ex)}"}
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass


def _validate_extract_excel(args):
    obj = args.get("object_name", "")
    if not obj or not isinstance(obj, str) or not obj.strip():
        return "缺少必填参数: object_name"
    if not (obj.lower().endswith(".xlsx") or obj.lower().endswith(".xls")):
        return f"不支持的文件格式，仅支持 .xlsx/.xls: {obj}"
    return None


@ToolRegistry.register(
    "extract_excel",
    "解析 Excel 文档 (.xlsx/.xls)，读取工作表数据。支持分页读取大表格。",
    {
        "type": "object",
        "properties": {
            "object_name": {
                "type": "string",
                "description": "MinIO 中的 Excel 对象路径 (object_name)，由文件上传接口返回",
            },
            "sheet_name": {"type": "string", "description": "要读取的工作表名称（可选，不传则读取第一个工作表）"},
            "row_offset": {"type": "integer", "description": "行起始偏移量（默认 0，用于分页读取大表格）"},
            "max_rows": {"type": "integer", "description": "最多读取的行数（默认 500，最大 2000）"},
        },
        "required": ["object_name"],
    },
    validator=_validate_extract_excel,
)
def _tool_extract_excel(args):
    object_name = args["object_name"].strip()
    sheet_name = (args.get("sheet_name") or "").strip() or None
    row_offset = max(0, args.get("row_offset", 0) or 0)
    max_rows = min(args.get("max_rows", 500) or 500, 2000)

    tmp_path = _download_from_minio(object_name)
    if not tmp_path:
        return {"success": False, "error": f"从 MinIO 下载文件失败: {object_name}"}

    try:
        is_legacy = object_name.lower().endswith(".xls") and not object_name.lower().endswith(".xlsx")

        if is_legacy:
            import xlrd

            wb = xlrd.open_workbook(tmp_path)
            sheet_names = wb.sheet_names()

            target = sheet_name or sheet_names[0]
            try:
                ws = wb.sheet_by_name(target)
            except xlrd.XLRDError:
                return {
                    "success": False,
                    "error": f"工作表 '{target}' 不存在，可用: {', '.join(sheet_names)}",
                }

            total_rows = ws.nrows
            data = []
            for ri in range(min(total_rows, row_offset + max_rows)):
                if ri < row_offset:
                    continue
                data.append(
                    [str(ws.cell_value(ri, ci)) if ws.cell_value(ri, ci) != "" else "" for ci in range(ws.ncols)]
                )
        else:
            from openpyxl import load_workbook

            wb = load_workbook(tmp_path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames

            target = sheet_name or sheet_names[0]
            if target not in sheet_names:
                return {
                    "success": False,
                    "error": f"工作表 '{target}' 不存在，可用: {', '.join(sheet_names)}",
                }

            ws = wb[target]
            data = []
            total_rows = 0
            for ri, row in enumerate(ws.iter_rows(values_only=True)):
                total_rows += 1
                if ri < row_offset:
                    continue
                if len(data) >= max_rows:
                    continue
                data.append([str(c) if c is not None else "" for c in row])

            wb.close()

        truncated = (row_offset + len(data)) < total_rows
        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "sheet_name": target,
                "sheet_names": sheet_names,
                "total_rows": total_rows,
                "row_offset": row_offset,
                "max_rows": max_rows,
                "rows_returned": len(data),
                "rows": data,
                "truncated": truncated,
                "hint": f"共 {total_rows} 行，当前返回第 {row_offset}–{row_offset + len(data)} 行。"
                + ("数据未读完，用 row_offset 参数继续。" if truncated else "已读完。"),
            },
        }
    except Exception as ex:
        logging.warning("extract_excel 解析失败: %s", ex)
        return {"success": False, "error": f"Excel 解析失败: {str(ex)}"}
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass


def _validate_read_text(args):
    obj = args.get("object_name", "")
    if not obj or not isinstance(obj, str) or not obj.strip():
        return "缺少必填参数: object_name"
    lower = obj.lower()
    if not (lower.endswith(".txt") or lower.endswith(".md")):
        return f"不支持的文件格式，仅支持 .txt/.md: {obj}"
    return None


@ToolRegistry.register(
    "read_text",
    "读取纯文本/Markdown 文件 (.txt/.md)，支持分页读取长文本。",
    {
        "type": "object",
        "properties": {
            "object_name": {"type": "string", "description": "MinIO 中的文本文件 object_name，由文件上传接口返回"},
            "char_offset": {"type": "integer", "description": "字符起始偏移量（默认 0，用于分页读取长文本）"},
            "char_limit": {"type": "integer", "description": "最多返回的字符数（默认 30000，最大 80000）"},
        },
        "required": ["object_name"],
    },
    validator=_validate_read_text,
)
def _tool_read_text(args):
    object_name = args["object_name"].strip()
    char_offset = max(0, args.get("char_offset", 0) or 0)
    char_limit = min(args.get("char_limit", 30000) or 30000, 80000)

    tmp_path = _download_from_minio(object_name)
    if not tmp_path:
        return {"success": False, "error": f"从 MinIO 下载文件失败: {object_name}"}

    def _read(encoding):
        with open(tmp_path, "r", encoding=encoding) as f:
            return f.read()

    try:
        content = _read("utf-8")
    except UnicodeDecodeError:
        try:
            content = _read("gbk")
        except Exception as ex:
            return {"success": False, "error": f"文件编码不支持 (UTF-8/GBK): {str(ex)}"}
    except Exception as ex:
        logging.warning("read_text 读取失败: %s", ex)
        return {"success": False, "error": f"文件读取失败: {str(ex)}"}
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass

    total = len(content)
    chunk = content[char_offset : char_offset + char_limit]
    truncated = (char_offset + char_limit) < total
    return {
        "success": True,
        "data": {
            "filename": os.path.basename(object_name),
            "content": chunk,
            "total_chars": total,
            "char_offset": char_offset,
            "char_limit": char_limit,
            "truncated": truncated,
            "hint": f"共 {total} 字符，当前返回第 {char_offset}–{char_offset + len(chunk)} 字符。"
            + ("内容未读完，用 char_offset 参数继续。" if truncated else "已读完。"),
        },
    }


def _validate_parse_json(args):
    obj = args.get("object_name", "")
    if not obj or not isinstance(obj, str) or not obj.strip():
        return "缺少必填参数: object_name"
    if not obj.lower().endswith(".json"):
        return f"不支持的文件格式，仅支持 .json: {obj}"
    return None


@ToolRegistry.register(
    "parse_json",
    "解析 JSON 文件 (.json)，返回结构化数据摘要。支持数组/对象分页、路径导航、大值截断。",
    {
        "type": "object",
        "properties": {
            "object_name": {"type": "string", "description": "MinIO 中的 JSON 文件 object_name，由文件上传接口返回"},
            "query_path": {
                "type": "string",
                "description": "可选的路径（如 'data.users'、'items[0].name'），用于提取深层字段",
            },
            "array_offset": {"type": "integer", "description": "数组元素起始索引（默认 0）"},
            "array_limit": {"type": "integer", "description": "数组最多返回元素数（默认 200，最大 2000）"},
            "key_offset": {"type": "integer", "description": "对象键起始偏移量（默认 0）"},
            "key_limit": {"type": "integer", "description": "对象最多返回键数（默认 50，最大 200）"},
            "max_list_items": {"type": "integer", "description": "嵌套列表最多返回项数（默认 100，最大 2000）"},
        },
        "required": ["object_name"],
    },
    validator=_validate_parse_json,
)
def _tool_parse_json(args):
    object_name = args["object_name"].strip()
    query_path = (args.get("query_path") or "").strip() or None
    array_offset = max(0, args.get("array_offset", 0) or 0)
    array_limit = min(args.get("array_limit", 200) or 200, 2000)
    key_offset = max(0, args.get("key_offset", 0) or 0)
    key_limit = min(args.get("key_limit", 50) or 50, 200)
    max_list = min(args.get("max_list_items", 100) or 100, 2000)

    MAX_VAL_LEN = 8000
    MAX_DICT_SAMPLE = 20
    MAX_DEPTH = 4

    def _summarize(val, depth=0, field_path=""):
        if depth > MAX_DEPTH:
            return f"<嵌套{type(val).__name__}>"
        if isinstance(val, str):
            if len(val) > MAX_VAL_LEN:
                return val[:MAX_VAL_LEN] + f"…(截断，共{len(val)}字符)"
            return val
        if isinstance(val, (int, float, bool)) or val is None:
            return val
        if isinstance(val, list):
            total = len(val)
            if total == 0:
                return []
            if total <= max_list:
                return [_summarize(v, depth + 1, f"{field_path}[{i}]") for i, v in enumerate(val)]
            sample = [_summarize(v, depth + 1, f"{field_path}[{i}]") for i, v in enumerate(val[:max_list])]
            return {
                "__head": sample,
                "__total": total,
                "__hint": f"数组共 {total} 项，仅返回前 {max_list} 项。用 query_path 定位到此字段后可用 array_offset/array_limit 分页读完。",
            }
        if isinstance(val, dict):
            ks = list(val.keys())
            total = len(ks)
            if total == 0:
                return {}
            if total <= MAX_DICT_SAMPLE:
                return {k: _summarize(v, depth + 1, f"{field_path}.{k}" if field_path else k) for k, v in val.items()}
            head = {
                k: _summarize(v, depth + 1, f"{field_path}.{k}" if field_path else k)
                for k, v in list(val.items())[:MAX_DICT_SAMPLE]
            }
            return {"__dict_head": head, "__total_keys": total}
        return str(val)[:500]

    tmp_path = _download_from_minio(object_name)
    if not tmp_path:
        return {"success": False, "error": f"从 MinIO 下载文件失败: {object_name}"}

    try:
        with open(tmp_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except UnicodeDecodeError:
        try:
            with open(tmp_path, "r", encoding="gbk") as f:
                data = json.load(f)
        except Exception as ex:
            return {"success": False, "error": f"JSON 解析失败 (编码): {str(ex)}"}
    except json.JSONDecodeError as ex:
        return {"success": False, "error": f"JSON 格式无效: {str(ex)}"}
    except Exception as ex:
        logging.warning("parse_json 读取失败: %s", ex)
        return {"success": False, "error": f"文件读取失败: {str(ex)}"}
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass

    # Path navigation
    target = data
    if query_path:
        try:
            for key in query_path.split("."):
                if "[" in key and key.endswith("]"):
                    name, idx = key[:-1].split("[", 1)
                    target = target[name][int(idx)]
                else:
                    target = target[key]
        except (KeyError, IndexError, TypeError, ValueError) as ex:
            return {"success": False, "error": f"路径 '{query_path}' 不可达: {str(ex)}"}

    if isinstance(target, list):
        total = len(target)
        page = target[array_offset : array_offset + array_limit]
        truncated = (array_offset + array_limit) < total
        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "type": "array",
                "total_items": total,
                "array_offset": array_offset,
                "array_limit": array_limit,
                "items_returned": len(page),
                "items": [_summarize(v) for v in page],
                "truncated": truncated,
                "hint": f"JSON 数组共 {total} 项，当前返回第 {array_offset}–{array_offset + len(page)} 项。"
                + ("未读完，用 array_offset 参数继续。" if truncated else "已读完。"),
            },
        }
    elif isinstance(target, dict):
        all_keys = list(target.keys())
        total_keys = len(all_keys)
        page_keys = all_keys[key_offset : key_offset + key_limit]
        page_items = {k: _summarize(target[k]) for k in page_keys}
        truncated = (key_offset + key_limit) < total_keys
        raw_size = len(json.dumps(target, ensure_ascii=False))
        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "type": "object",
                "total_keys": total_keys,
                "key_offset": key_offset,
                "key_limit": key_limit,
                "keys_returned": len(page_keys),
                "items": page_items,
                "raw_size": raw_size,
                "truncated": truncated,
                "hint": f"JSON 对象共 {total_keys} 个键（{raw_size} 字符），当前返回第 {key_offset}–{key_offset + len(page_keys)} 个键。"
                + ("未读完，用 key_offset 参数继续。" if truncated else "已读完。"),
            },
        }
    else:
        return {
            "success": True,
            "data": {
                "filename": os.path.basename(object_name),
                "type": type(target).__name__,
                "value": _summarize(target),
            },
        }
